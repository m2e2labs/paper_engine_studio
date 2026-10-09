"""
make.py  -  the sample workbook for Excel for Absolute Beginners, and its answers

One small made-up shop, one year of sales (2025), with four sheets:
Sales (one row per sale), Products, Targets and Staff. The data is clean on
purpose: beginners should meet the features first, and the cleaning comes later.

Everything is deterministic (one seed), so every figure printed on a practice page
can be regenerated and checked. answers.json holds every figure a page prints.

    python datasets/excel-beginner/make.py      writes into this folder
"""
import json, os, random
from datetime import date, timedelta
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SEED = 11
rng = random.Random(SEED)

PRODUCTS = [  # name, category, unit price
    ("Notebook", "Stationery", 3),
    ("Pen", "Stationery", 2),
    ("Mug", "Kitchen", 8),
    ("Bag", "Kitchen", 12),
    ("Plant", "Garden", 15),
]
STAFF = [  # name, region, start date
    ("Ada Okafor", "North", date(2019, 3, 4)),
    ("Ben Lindqvist", "South", date(2021, 6, 1)),
    ("Carla Moreau", "East", date(2022, 9, 12)),
    ("Dev Patel", "West", date(2023, 1, 16)),
    ("Elin Haddad", "North", date(2024, 2, 5)),
]
TARGETS = [("North", 1500), ("South", 1400), ("East", 1300), ("West", 1200)]
REGIONS = [r for r, _ in TARGETS]
YEAR_START, YEAR_END = date(2025, 1, 1), date(2025, 12, 31)
N_SALES = 240

# Sales: a random day in 2025, a salesperson (so the region follows from them),
# a product, and a quantity from 1 to 6. The amount is quantity times the unit price.
prices = {name: price for name, _, price in PRODUCTS}
category = {name: cat for name, cat, _ in PRODUCTS}
span = (YEAR_END - YEAR_START).days
sales = []
for _ in range(N_SALES):
    d = YEAR_START + timedelta(days=rng.randrange(span + 1))
    person, region, _ = rng.choice(STAFF)
    product = rng.choice([p for p, _, _ in PRODUCTS])
    qty = rng.randint(1, 6)
    sales.append({"date": d, "region": region, "person": person,
                  "product": product, "qty": qty, "price": prices[product],
                  "amount": qty * prices[product]})
sales.sort(key=lambda s: s["date"])

# ---------------------------------------------------------------- the workbook
HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="EEF0FE")

def sheet(wb, title, headers, rows, widths, date_cols=(), money_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for col in date_cols:
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col).number_format = "dd/mm/yyyy"
    for col in money_cols:
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col).number_format = "£#,##0.00"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws

wb = Workbook()
wb.remove(wb.active)
sheet(wb, "Sales",
      ["Date", "Region", "Salesperson", "Product", "Quantity", "Unit price", "Amount"],
      [[s["date"], s["region"], s["person"], s["product"], s["qty"], s["price"], s["amount"]] for s in sales],
      [12, 10, 18, 12, 10, 12, 12], date_cols=(1,), money_cols=(6, 7))
sheet(wb, "Products", ["Product", "Category", "Unit price"],
      [[n, c, p] for n, c, p in PRODUCTS], [14, 14, 12], money_cols=(3,))
sheet(wb, "Targets", ["Region", "Target"],
      [[r, t] for r, t in TARGETS], [12, 12])
sheet(wb, "Staff", ["Name", "Region", "Start date"],
      [[n, r, d] for n, r, d in STAFF], [18, 12, 14], date_cols=(3,))

out = os.path.join(HERE, "Corner Shop sales.xlsx")
wb.save(out)

# ---------------------------------------------------------------- the answers
amounts = [s["amount"] for s in sales]
by_region = defaultdict(int)
by_product = defaultdict(int)
by_category = defaultdict(int)
count_region = defaultdict(int)
for s in sales:
    by_region[s["region"]] += s["amount"]
    by_product[s["product"]] += s["amount"]
    by_category[category[s["product"]]] += s["amount"]
    count_region[s["region"]] += 1

answers = {
    "rows": len(sales),
    "total_amount": sum(amounts),
    "average_amount": round(sum(amounts) / len(amounts), 2),
    "min_amount": min(amounts),
    "max_amount": max(amounts),
    "count_amount_over_50": sum(1 for a in amounts if a > 50),
    "rows_per_region": dict(sorted(count_region.items())),
    "amount_by_region": dict(sorted(by_region.items())),
    "amount_by_product": dict(sorted(by_product.items())),
    "amount_by_category": dict(sorted(by_category.items())),
    "north_total": by_region["North"],
    "bag_rows": sum(1 for s in sales if s["product"] == "Bag"),
    "first_sale": sales[0]["date"].isoformat(),
    "last_sale": sales[-1]["date"].isoformat(),
}
with open(os.path.join(HERE, "answers.json"), "w", encoding="utf8") as fh:
    json.dump(answers, fh, indent=2)

print("wrote", out, "and answers.json:", answers["rows"], "rows, total", answers["total_amount"])
