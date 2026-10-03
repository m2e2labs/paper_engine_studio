"""
make.py  -  the sample files for the Absolute Beginners books, and their answers

One small bike shop, two years of sales, written the way data really arrives:
a folder of monthly workbooks with a title row above the headers, one sheet laid
out for a person to read (a column per month), a Products sheet that doesn't
quite match, and a clean workbook for the books that start after the cleaning.

Everything is deterministic (one seed), so the figures printed on the practice
pages can be regenerated from scratch and checked. answers.json holds every
figure a page prints.

    python datasets/absolute-beginners/make.py      writes into this folder
"""
import csv, json, os, random
from datetime import date, timedelta
from collections import defaultdict
from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SEED = 7
rng = random.Random(SEED)

REGIONS = ["North", "South", "East", "West"]
# Products the shop sells. Lamp is sold but missing from the Products sheet, and
# Pump is on the sheet but never sold: both on purpose, for the merge lessons.
PRODUCTS = {  # name: (category, colour, list price, the prices it actually sold at)
    "Bike":   ("Cycles", "Blue",  420, [420, 420, 420, 380]),
    "Helmet": ("Gear",   "Black",  45, [45, 45, 39]),
    "Lock":   ("Gear",   "Black",  18, [18, 18, 15]),
    "Bell":   ("Gear",   "Red",     6, [6]),
    "Pump":   ("Gear",   "Red",    15, [15]),
}
SOLD = ["Bike", "Helmet", "Lock", "Lamp", "Bell"]
LAMP_PRICES = [22, 22, 19]
WEIGHTS = {"Bike": 3, "Helmet": 5, "Lock": 6, "Lamp": 4, "Bell": 5}

FIRST = ["Ada", "Ben", "Carla", "Dev", "Elin", "Femi", "Grace", "Hugo", "Iris", "Jonas",
         "Kira", "Leo", "Mara", "Nils", "Omar", "Priya", "Quinn", "Rosa", "Sam", "Tomas"]
LAST = ["Okafor", "Lindqvist", "Moreau", "Patel", "Haddad", "Novak", "Ibarra", "Kowalski",
        "Tanaka", "Mensah", "Berg", "Costa", "Doyle", "Eriksen", "Farah", "Gupta",
        "Holm", "Iqbal", "Jensen", "Kaur"]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
LONG = ["January", "February", "March", "April", "May", "June", "July", "August",
        "September", "October", "November", "December"]

# ---------------------------------------------------------------- customers
customers = []
for i in range(20):
    customers.append({"CustomerID": f"C{i + 1:02d}", "Name": f"{FIRST[i]} {LAST[i]}", "Region": REGIONS[i % 4]})
region_of = {c["CustomerID"]: c["Region"] for c in customers}

# ---------------------------------------------------------------- order lines
lines = []
for year in (2024, 2025):
    for m in range(1, 13):
        n = rng.randint(42, 58)
        days = (date(year + (m == 12), m % 12 + 1, 1) - date(year, m, 1)).days
        for _ in range(n):
            c = rng.choice(customers)
            p = rng.choices(SOLD, weights=[WEIGHTS[s] for s in SOLD])[0]
            price = rng.choice(LAMP_PRICES if p == "Lamp" else PRODUCTS[p][3])
            qty = rng.choice([1, 1, 1, 2, 2, 3]) if p != "Bike" else rng.choice([1, 1, 1, 2])
            lines.append({
                "OrderDate": date(year, m, rng.randint(1, days)), "Region": c["Region"],
                "CustomerID": c["CustomerID"], "Product": p, "Qty": qty, "Price": price, "Amount": qty * price,
            })
lines.sort(key=lambda r: (r["OrderDate"], r["CustomerID"]))

# the dirt, on purpose: three lines with no quantity (an order that was never filled in),
# and four February 2025 lines whose Region arrived with a space in front of it.
blank_qty = [i for i, r in enumerate(lines) if r["OrderDate"].year == 2025 and r["Product"] == "Lock"][:3]
for i in blank_qty:
    lines[i]["Qty"] = None; lines[i]["Amount"] = None
feb = [i for i, r in enumerate(lines) if r["OrderDate"].year == 2025 and r["OrderDate"].month == 2 and r["Region"] == "North"][:4]
RAW_REGION = {i: " North" for i in feb}

# ---------------------------------------------------------------- writers
def sheet(wb, name, title, header, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    if title: ws.append([title])
    ws.append(header)
    for r in rows: ws.append(r)
    return ws

os.makedirs(os.path.join(HERE, "Sales"), exist_ok=True)
for year in (2024, 2025):
    for m in range(1, 13):
        wb = Workbook()
        rows = [[r["OrderDate"], RAW_REGION.get(i, r["Region"]), r["CustomerID"], r["Product"], r["Qty"], r["Price"], r["Amount"]]
                for i, r in enumerate(lines) if r["OrderDate"].year == year and r["OrderDate"].month == m]
        ws = sheet(wb, "Sales", f"Sales, {LONG[m - 1]} {year}", ["OrderDate", "Region", "CustomerID", "Product", "Qty", "Price", "Amount"], rows, first=True)
        for row in ws.iter_rows(min_row=3, min_col=1, max_col=1):
            row[0].number_format = "yyyy-mm-dd"
        wb.save(os.path.join(HERE, "Sales", f"Sales {year}-{m:02d}.xlsx"))

# the sheet built for people: 2025 units, a column per month, a title row on top
wide = defaultdict(lambda: [0] * 12)
for r in lines:
    if r["OrderDate"].year == 2025 and r["Qty"]:
        wide[(r["Region"], r["Product"])][r["OrderDate"].month - 1] += r["Qty"]
wide_rows = [[reg, p] + wide[(reg, p)] for reg in REGIONS for p in SOLD]
products_rows = [[p, v[0], v[1], v[2]] for p, v in PRODUCTS.items()]
wb = Workbook()
sheet(wb, "Sales", "Sales by month", ["Region", "Product"] + MONTHS, wide_rows, first=True)
sheet(wb, "Products", None, ["Product", "Category", "Colour", "Price"], products_rows)
wb.save(os.path.join(HERE, "Sales by month.xlsx"))

# the clean workbook, for the books that start after the cleaning
wb = Workbook()
ws = sheet(wb, "Sales", None, ["OrderDate", "Region", "CustomerID", "Product", "Qty", "Price", "Amount"],
           [[r["OrderDate"], r["Region"], r["CustomerID"], r["Product"], r["Qty"], r["Price"], r["Amount"]] for r in lines], first=True)
for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
    row[0].number_format = "yyyy-mm-dd"
sheet(wb, "Products", None, ["Product", "Category", "Colour", "Price"], products_rows)
sheet(wb, "Customers", None, ["CustomerID", "Name", "Region"], [[c["CustomerID"], c["Name"], c["Region"]] for c in customers])
wb.save(os.path.join(HERE, "Sales clean.xlsx"))

with open(os.path.join(HERE, "Customers.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["CustomerID", "Name", "Region"])
    for c in customers: w.writerow([c["CustomerID"], c["Name"], c["Region"]])
targets = {"North": 24000, "South": 21000, "East": 18000, "West": 16000}
with open(os.path.join(HERE, "Targets.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["Region", "Target"])
    for reg, t in targets.items(): w.writerow([reg, t])

# ---------------------------------------------------------------- the answers
def total(rows): return sum(r["Amount"] or 0 for r in rows)
def qty(rows): return sum(r["Qty"] or 0 for r in rows)
y25 = [r for r in lines if r["OrderDate"].year == 2025]
y24 = [r for r in lines if r["OrderDate"].year == 2024]
by = lambda rows, key: {k: total([r for r in rows if r[key] == k]) for k in sorted({r[key] for r in rows})}
cat_of = lambda p: PRODUCTS[p][0] if p in PRODUCTS else None
feb25 = [r for r in lines if r["OrderDate"].year == 2025 and r["OrderDate"].month == 2]
mar25 = [r for r in lines if r["OrderDate"].year == 2025 and r["OrderDate"].month == 3]
jan25 = [r for r in lines if r["OrderDate"].year == 2025 and r["OrderDate"].month == 1]
h1 = [r for r in y25 if r["OrderDate"].month <= 6]
prices = [r["Price"] for r in lines]
blue = [r for r in lines if r["Product"] == "Bike"]
per_customer = {c["CustomerID"]: total([r for r in lines if r["CustomerID"] == c["CustomerID"]]) for c in customers}
north_rows = [r for r in lines if r["Region"] == "North"]
answers = {
    "seed": SEED,
    "files": {"monthly_workbooks": 24, "products_on_sheet": len(PRODUCTS), "products_sold": len(SOLD), "customers": len(customers),
              "month_columns": 12, "wide_rows": len(wide_rows)},
    "lines": {"all": len(lines), "2024": len(y24), "2025": len(y25), "jan_2025": len(jan25), "feb_2025": len(feb25), "mar_2025": len(mar25),
              "feb_plus_mar_2025": len(feb25) + len(mar25), "blank_qty": len(blank_qty), "leading_space_regions": len(feb),
              "north": len(north_rows), "feb_north_distinct_raw": len({RAW_REGION.get(i, r["Region"]) for i, r in enumerate(lines) if r in feb25 and r["Region"] == "North"})},
    "total_sales": {"all": total(lines), "2024": total(y24), "2025": total(y25), "h1_2025": total(h1), "jan_2025": total(jan25),
                    "by_region": by(lines, "Region"), "by_product": by(lines, "Product"),
                    "by_category": {"Cycles": total([r for r in lines if cat_of(r["Product"]) == "Cycles"]),
                                    "Gear": total([r for r in lines if cat_of(r["Product"]) == "Gear"]),
                                    "blank": total([r for r in lines if cat_of(r["Product"]) is None])},
                    "west": total([r for r in lines if r["Region"] == "West"]), "blue": total(blue),
                    "by_region_2025": by(y25, "Region")},
    "units": {"all": qty(lines), "2025": qty(y25), "jan_2025_wide": sum(v[0] for v in wide.values()),
              "unpivoted_rows": len(wide_rows) * 12, "q1_north_bike": sum(wide[("North", "Bike")][:3]),
              "by_month_2025": {MONTHS[i]: sum(v[i] for v in wide.values()) for i in range(12)}},
    "aggregates": {"average_amount": round(total([r for r in lines if r["Amount"]]) / len([r for r in lines if r["Amount"]]), 2),
                   "min_amount": min(r["Amount"] for r in lines if r["Amount"]), "max_amount": max(r["Amount"] for r in lines if r["Amount"]),
                   "count_amount": len([r for r in lines if r["Amount"] is not None]), "countrows": len(lines),
                   "distinct_products": len(SOLD), "sum_qty": qty(lines), "average_price": round(sum(prices) / len(prices), 2),
                   "wrong_revenue": round(qty(lines) * sum(prices) / len(prices), 2), "sumx_revenue": total(lines),
                   "averagex_customer": round(sum(per_customer.values()) / len(per_customer), 2),
                   "top_customer": max(per_customer, key=per_customer.get), "top_customer_sales": max(per_customer.values()),
                   "pct_cycles": round(100 * total([r for r in lines if cat_of(r["Product"]) == "Cycles"]) / total(lines), 1),
                   "bands": {"small (under 50)": len([r for r in lines if r["Amount"] is not None and r["Amount"] < 50]),
                             "medium (50 to 399)": len([r for r in lines if r["Amount"] is not None and 50 <= r["Amount"] < 400]),
                             "large (400 and up)": len([r for r in lines if r["Amount"] is not None and r["Amount"] >= 400])},
                   "date_rows": (date(2025, 12, 31) - date(2024, 1, 1)).days + 1,
                   "ytd_jun_2025": total(h1), "ly_jun_2025": total([r for r in y24 if r["OrderDate"].month <= 6]),
                   "jun_2025": total([r for r in y25 if r["OrderDate"].month == 6]), "jun_2024": total([r for r in y24 if r["OrderDate"].month == 6]),
                   "top_region": max(by(lines, "Region"), key=by(lines, "Region").get), "top_product": max(by(lines, "Product"), key=by(lines, "Product").get)},
    "targets": targets,
    "merge": {"sales_product_without_match": "Lamp", "products_row_without_sales": "Pump",
              "left_outer_rows": len(lines), "inner_rows": len([r for r in lines if r["Product"] in PRODUCTS]),
              "left_anti_rows": len([r for r in lines if r["Product"] not in PRODUCTS])},
}
with open(os.path.join(HERE, "answers.json"), "w", encoding="utf-8") as f:
    json.dump(answers, f, indent=2, default=str)
print(json.dumps(answers, indent=1, default=str))
