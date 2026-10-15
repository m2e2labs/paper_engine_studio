"""
make.py  -  the messy sample files for Excel Power Query for Absolute Beginners, and their answers

The same 240 Corner Shop sales as Excel for Absolute Beginners, written the way data really
arrives, so every cleaned total still matches datasets/excel-beginner/answers.json (240 rows,
6325 in all). The sales are READ from datasets/excel-beginner/Corner Shop sales.xlsx, never
made up again, so the two books can't drift apart.

    Corner Shop export.xlsx      Export (the messy export), List (a lightly untidy list for a
                                 first query), Products (one product missing, one never sold),
                                 Staff list (text to split and merge), Targets sideways
    Corner Shop sales.csv        the 240 sales as a CSV
    Monthly/                     12 tidy monthly workbooks, headings in the first row
    Monthly messy/               the same 12 months with a title row and a Total row in each
    Sales by month report.xlsx   Wide (months across the top) and Report (built for people:
                                 titles, a region shown once per block, subtotals)
    Year-end export.csv          a fresh messy file for the last practice page

The dirt is deliberate and deterministic (one seed). answers.json holds every figure a
practice page can print.

    python datasets/excel-power-query/make.py      writes into this folder
"""
import csv, json, os, random, shutil
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE = os.path.join(HERE, "..", "excel-beginner", "Corner Shop sales.xlsx")
SEED = 23
rng = random.Random(SEED)

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
LONG = ["January", "February", "March", "April", "May", "June", "July", "August",
        "September", "October", "November", "December"]
HEADERS = ["Date", "Region", "Salesperson", "Product", "Quantity", "Unit price", "Amount"]
BOLD = Font(bold=True)
FILL = PatternFill("solid", fgColor="EEF0FE")

# ---------------------------------------------------------------- the 240 sales
src = load_workbook(SOURCE)
sales = []
for d, region, person, product, qty, price, amount in src["Sales"].iter_rows(min_row=2, values_only=True):
    sales.append({"date": d.date(), "region": region, "person": person, "product": product,
                  "qty": qty, "price": price, "amount": amount})
staff = [r for r in src["Staff"].iter_rows(min_row=2, values_only=True)]
targets = [r for r in src["Targets"].iter_rows(min_row=2, values_only=True)]
category = {n: c for n, c, _ in src["Products"].iter_rows(min_row=2, values_only=True)}
for i, s in enumerate(sales, start=1):
    s["id"] = f"S{i:04d}"
assert len(sales) == 240 and sum(s["amount"] for s in sales) == 6325, "the source workbook changed"

uk = lambda d: d.strftime("%d/%m/%Y")

def widths(ws, cols):
    for i, w in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def header_row(ws, row):
    for cell in ws[row]:
        if cell.value is not None:
            cell.font, cell.fill = BOLD, FILL

def clean_sheet(ws, rows):
    ws.append(HEADERS)
    header_row(ws, 1)
    for s in rows:
        ws.append([s["date"], s["region"], s["person"], s["product"], s["qty"], s["price"], s["amount"]])
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
    widths(ws, [12, 10, 18, 12, 10, 12, 10])

answers = {"source": "datasets/excel-beginner/Corner Shop sales.xlsx", "rows": 240, "total_amount": 6325}

# ---------------------------------------------------------------- Corner Shop export.xlsx : Export
# Headings in row 4 under a title, the export date and a blank row. Dates are text, day first.
# Region is left empty when it repeats the row above (the till system "suppresses" repeats), and
# some of the Regions that do show have stray spaces or a misspelling. Some Products have the
# wrong capitals. Amount is text, and a few amounts carry a pound sign, so changing the column
# to a number gives errors until the sign is replaced. Four sales are exported twice, five blank
# rows sit in the data, and the export ends with a blank row and an "End of export" line.
pick = lambda n, pool: set(rng.sample(sorted(pool), n))
dupes = pick(4, range(240))
pound = pick(6, set(range(240)) - dupes)
product_case = pick(18, range(240))

order = []
for i in range(240):
    order.append(i)
    if i in dupes:
        order.append(i)                      # the same sale, exported twice in a row
data_rows = []                               # (sale index, region shown)
prev = None
for i in order:
    shown = None if sales[i]["region"] == prev else sales[i]["region"]
    prev = sales[i]["region"]
    data_rows.append([i, shown])
visible = [k for k, (_, shown) in enumerate(data_rows) if shown is not None]
misspelt = pick(3, [k for k in visible if data_rows[k][1] == "North"])
spaced = pick(12, set(visible) - misspelt)
for k in misspelt:
    data_rows[k][1] = "Nrth"
for n, k in enumerate(sorted(spaced)):
    data_rows[k][1] = (" " + data_rows[k][1]) if n % 2 == 0 else (data_rows[k][1] + " ")

def export_values(i, shown):
    s = sales[i]
    product = s["product"]
    if i in product_case:
        product = product.upper() if i % 2 else product.lower()
    amount = ("£" if i in pound else "") + str(s["amount"])
    return [s["id"], uk(s["date"]), shown, s["person"], product, s["qty"], s["price"], amount,
            f"Till {1 + i % 2}", "system"]

blank_at = sorted(rng.sample(range(10, len(data_rows) - 10), 5))
wb = Workbook()
ws = wb.active
ws.title = "Export"
ws.append(["Corner Shop sales export"])
ws.append(["Exported 05/01/2026 from the till system"])
ws.append([])
ws.append(["Sale ID", "Sale date", "Region", "Salesperson", "Product", "Qty", "Unit price", "Amount", "Till", "Exported by"])
header_row(ws, 4)
ws["A1"].font = Font(bold=True, size=14)
for k, (i, shown) in enumerate(data_rows):
    if k in blank_at:
        ws.append([])
    ws.append(export_values(i, shown))
ws.append([])
ws.append([f"End of export, {len(data_rows)} rows"])
widths(ws, [10, 12, 10, 18, 12, 6, 10, 10, 8, 12])
answers["export"] = {
    "sheet": "Export",
    "heading_row": 4,
    "rows_above_headings": 3,
    "rows_after_promoting_headings": ws.max_row - 4,
    "data_rows_with_duplicates": len(data_rows),
    "blank_rows_in_data": 5,
    "blank_rows_at_the_end": 1,
    "end_line": f"End of export, {len(data_rows)} rows",
    "duplicate_sale_ids": sorted(sales[i]["id"] for i in dupes),
    "region_cells_left_empty": sum(1 for _, shown in data_rows if shown is None),
    "region_cells_with_stray_spaces": len(spaced),
    "region_cells_misspelt_Nrth": len(misspelt),
    "product_cells_with_wrong_capitals": len(product_case),
    "amount_cells_with_a_pound_sign": len(pound),
    "pound_sign_sale_ids": sorted(sales[i]["id"] for i in pound),
    "amount_lost_if_errors_are_removed_not_fixed": sum(sales[i]["amount"] for i in pound),
    "rows_if_errors_are_removed_not_fixed": 240 - len(pound),
    "total_if_errors_are_removed_not_fixed": 6325 - sum(sales[i]["amount"] for i in pound),
    "rows_after_cleaning": 240,
    "total_after_cleaning": 6325,
}

# ---------------------------------------------------------------- : List
# A lightly untidy list starting in A1, for the first query: plain range, real dates and
# numbers, but a few stray spaces in Region and some lower-case products.
ws = wb.create_sheet("List")
list_spaced = pick(10, range(240))
list_lower = pick(10, set(range(240)) - list_spaced)
ws.append(HEADERS)
header_row(ws, 1)
for i, s in enumerate(sales):
    region = s["region"] + " " if i in list_spaced else s["region"]
    product = s["product"].lower() if i in list_lower else s["product"]
    ws.append([s["date"], region, s["person"], product, s["qty"], s["price"], s["amount"]])
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
widths(ws, [12, 10, 18, 12, 10, 12, 10])
answers["list"] = {"sheet": "List", "rows": 240, "total_amount": 6325,
                   "region_cells_with_a_trailing_space": len(list_spaced),
                   "product_cells_in_lower_case": len(list_lower)}

# ---------------------------------------------------------------- : Products
# Plant is sold but missing from this list, and Candle is on the list but never sold, so the
# join kinds give different row counts.
ws = wb.create_sheet("Products")
ws.append(["Product", "Category", "Unit price"])
header_row(ws, 1)
for row in [("Notebook", "Stationery", 3), ("Pen", "Stationery", 2), ("Mug", "Kitchen", 8),
            ("Bag", "Kitchen", 12), ("Candle", "Home", 6)]:
    ws.append(list(row))
widths(ws, [12, 12, 10])

# ---------------------------------------------------------------- : Staff list
# Text to split, extract and merge: a team label with a dash, a fixed-width staff code
# (three letters of region, the start year, two initials), and first and last names apart.
ws = wb.create_sheet("Staff list")
ws.append(["Team", "Staff code", "First name", "Last name"])
header_row(ws, 1)
staff_rows = []
for name, region, start in staff:
    first, last = name.split(" ", 1)
    code = f"{region[:3].upper()}{start.year}{first[0]}{last[0]}"
    staff_rows.append({"team": f"{region} - {name}", "code": code, "first": first, "last": last})
    ws.append([f"{region} - {name}", code, first, last])
widths(ws, [24, 12, 12, 14])
answers["staff_list"] = {"sheet": "Staff list", "rows": staff_rows}

# ---------------------------------------------------------------- : Targets sideways
ws = wb.create_sheet("Targets sideways")
ws.append(["Region"] + [r for r, _ in targets])
ws.append(["Target"] + [t for _, t in targets])
ws["A1"].font = ws["A2"].font = BOLD
widths(ws, [10, 10, 10, 10, 10])
answers["targets_sideways"] = {"sheet": "Targets sideways", "rows": 2, "columns": 5,
                               "after_transpose_and_headings": [list(t) for t in targets]}
wb.save(os.path.join(HERE, "Corner Shop export.xlsx"))

# ---------------------------------------------------------------- Corner Shop sales.csv
with open(os.path.join(HERE, "Corner Shop sales.csv"), "w", newline="", encoding="utf-8") as fh:
    w = csv.writer(fh)
    w.writerow(HEADERS)
    for s in sales:
        w.writerow([uk(s["date"]), s["region"], s["person"], s["product"], s["qty"], s["price"], s["amount"]])
answers["csv"] = {"file": "Corner Shop sales.csv", "rows": 240, "total_amount": 6325,
                  "date_format": "dd/mm/yyyy, as text"}

# ---------------------------------------------------------------- Monthly/ and Monthly messy/
by_month = defaultdict(list)
for s in sales:
    by_month[s["date"].month].append(s)
for folder in ("Monthly", "Monthly messy"):
    path = os.path.join(HERE, folder)
    if os.path.isdir(path):
        shutil.rmtree(path)
    os.makedirs(path)
month_answers = {}
messy_spaced_total = 0
for m in range(1, 13):
    rows = by_month[m]
    name = f"Sales 2025-{m:02d}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    clean_sheet(ws, rows)
    wb.save(os.path.join(HERE, "Monthly", name))

    # the messy month: a title row and a blank row above the headings, a Total row at the
    # bottom, and a couple of Regions with a stray space
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append([f"Corner Shop sales, {LONG[m - 1]} 2025"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(HEADERS)
    header_row(ws, 3)
    spaced_here = set(rng.sample(range(len(rows)), min(2, len(rows))))
    messy_spaced_total += len(spaced_here)
    for k, s in enumerate(rows):
        region = " " + s["region"] if k in spaced_here else s["region"]
        ws.append([s["date"], region, s["person"], s["product"], s["qty"], s["price"], s["amount"]])
    ws.append(["Total", None, None, None, None, None, sum(s["amount"] for s in rows)])
    ws.cell(row=ws.max_row, column=1).font = ws.cell(row=ws.max_row, column=7).font = BOLD
    for r in range(4, ws.max_row):
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
    widths(ws, [12, 10, 18, 12, 10, 12, 10])
    wb.save(os.path.join(HERE, "Monthly messy", name))
    month_answers[f"2025-{m:02d}"] = {"rows": len(rows), "amount": sum(s["amount"] for s in rows)}
answers["monthly"] = {"folder": "Monthly", "files": 12, "rows": 240, "total_amount": 6325,
                      "by_month": month_answers}
answers["monthly_messy"] = {"folder": "Monthly messy", "files": 12, "rows_above_headings": 2,
                            "total_rows": 12, "region_cells_with_a_leading_space": messy_spaced_total,
                            "rows_after_cleaning": 240, "total_after_cleaning": 6325,
                            "rows_if_total_rows_are_kept": 252,
                            "amount_if_total_rows_are_kept": 6325 * 2}

# Combining two lists: January and February appended, then merged with the Products sheet
# above (Plant missing, Candle never sold).
jf = by_month[1] + by_month[2]
plant = [s for s in jf if s["product"] == "Plant"]
answers["jan_feb_with_products"] = {
    "january_rows": len(by_month[1]), "february_rows": len(by_month[2]),
    "appended_rows": len(jf), "appended_amount": sum(s["amount"] for s in jf),
    "plant_rows": len(plant), "plant_amount": sum(s["amount"] for s in plant),
    "left_outer_rows": len(jf), "left_outer_rows_with_no_category": len(plant),
    "inner_rows": len(jf) - len(plant), "inner_amount": sum(s["amount"] for s in jf) - sum(s["amount"] for s in plant),
    "left_anti_rows": len(plant),
    "amount_by_category_left_outer": dict(sorted(
        {c: sum(s["amount"] for s in jf if s["product"] != "Plant" and category[s["product"]] == c)
         for c in ("Stationery", "Kitchen")}.items())),
}

# ---------------------------------------------------------------- Sales by month report.xlsx
cell = defaultdict(int)                       # (region, product, month) -> amount
for s in sales:
    cell[(s["region"], s["product"], s["date"].month)] += s["amount"]
REGIONS = ["North", "South", "East", "West"]
PRODUCTS = ["Notebook", "Pen", "Mug", "Bag", "Plant"]
wb = Workbook()

# Wide: one row per region, a column per month (headings like "Jan 2025").
ws = wb.active
ws.title = "Wide"
ws.append(["Region"] + [f"{m} 2025" for m in MONTHS])
header_row(ws, 1)
region_month = {(r, m): sum(cell[(r, p, m)] for p in PRODUCTS) for r in REGIONS for m in range(1, 13)}
for r in REGIONS:
    ws.append([r] + [region_month[(r, m)] for m in range(1, 13)])
widths(ws, [10] + [10] * 12)
answers["wide"] = {
    "sheet": "Wide", "rows": 4, "month_columns": 12, "rows_after_unpivot": 48,
    "month_headings": [f"{m} 2025" for m in MONTHS],
    "amount_by_month": {MONTHS[m - 1]: sum(region_month[(r, m)] for r in REGIONS) for m in range(1, 13)},
    "amount_by_region": {r: sum(region_month[(r, m)] for m in range(1, 13)) for r in REGIONS},
    "total_amount": sum(region_month.values()),
}

# Report: built for a person to read. Two title rows and a blank row, then headings; each
# region is a block with the region named only on its first row, a subtotal row and a blank
# row after it; a Grand total at the end. Months are "Jan" to "Dec", and a month with no
# sales is left empty.
ws = wb.create_sheet("Report")
ws.append(["Corner Shop"])
ws.append(["Sales by region and product, 2025"])
ws.append([])
ws.append(["Region", "Product"] + MONTHS + ["Total"])
ws["A1"].font = Font(bold=True, size=14)
header_row(ws, 4)
report_cells = 0
for r in REGIONS:
    for k, p in enumerate(PRODUCTS):
        vals = [cell[(r, p, m)] or None for m in range(1, 13)]
        report_cells += sum(1 for v in vals if v)
        ws.append([r if k == 0 else None, p] + vals + [sum(v or 0 for v in vals)])
    sub = [sum(cell[(r, p, m)] for p in PRODUCTS) for m in range(1, 13)]
    ws.append([f"{r} total", None] + sub + [sum(sub)])
    for c in ws[ws.max_row]:
        c.font = BOLD
    ws.append([])
grand = [sum(cell[(r, p, m)] for r in REGIONS for p in PRODUCTS) for m in range(1, 13)]
ws.append(["Grand total", None] + grand + [sum(grand)])
for c in ws[ws.max_row]:
    c.font = BOLD
widths(ws, [13, 11] + [7] * 12 + [8])
answers["report"] = {
    "sheet": "Report", "rows_above_headings": 3, "product_rows": 20,
    "subtotal_rows": 4, "blank_rows": 4, "grand_total_rows": 1,
    "rows_after_promoting_headings": ws.max_row - 4,
    "rows_after_removing_junk": 20,
    "rows_after_unpivot": report_cells,
    "empty_month_cells": 20 * 12 - report_cells,
    "subtotals": {r: sum(region_month[(r, m)] for m in range(1, 13)) for r in REGIONS},
    "grand_total": sum(grand),
}
wb.save(os.path.join(HERE, "Sales by month report.xlsx"))

# ---------------------------------------------------------------- Year-end export.csv
# A fresh messy file for the last practice page, with different dirt from the Export sheet:
# semicolons between values, two title lines and a blank line, Region in mixed capitals,
# Amount with a pound sign on every value, three sales twice, and a Total line at the end.
ye_dupes = pick(3, range(240))
mixed = pick(30, range(240))
lines = 0
with open(os.path.join(HERE, "Year-end export.csv"), "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh, delimiter=";")
    w.writerow(["Corner Shop year-end export"]); w.writerow(["Generated 02/01/2026"]); w.writerow([])
    w.writerow(["Sale ID", "Date", "Region", "Salesperson", "Product", "Qty", "Price", "Amount"])
    for i, s in enumerate(sales):
        region = (s["region"].upper() if i % 2 else s["region"].lower()) if i in mixed else s["region"]
        row = [s["id"], uk(s["date"]), region, s["person"], s["product"], s["qty"], s["price"], f"£{s['amount']}"]
        for _ in range(2 if i in ye_dupes else 1):
            w.writerow(row); lines += 1
    w.writerow(["Total", "", "", "", "", "", "", "£6325"])
answers["year_end_csv"] = {
    "file": "Year-end export.csv", "delimiter": "semicolon", "encoding": "UTF-8 with a byte order mark",
    "lines_above_headings": 3, "sale_lines_with_duplicates": lines,
    "duplicate_sale_ids": sorted(sales[i]["id"] for i in ye_dupes),
    "region_cells_in_wrong_capitals": len(mixed), "total_line": "Total ... £6325",
    "rows_after_cleaning": 240, "total_after_cleaning": 6325,
}

# ---------------------------------------------------------------- answers.json
with open(os.path.join(HERE, "answers.json"), "w", encoding="utf8") as fh:
    json.dump(answers, fh, indent=2, ensure_ascii=False, default=str)
print("wrote the Excel Power Query sample files: 240 sales, total 6325, in every file")
